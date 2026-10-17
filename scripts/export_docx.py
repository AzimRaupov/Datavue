#!/usr/bin/env python3
"""Собирает Word-документ из картинок виджетов (их снимает export-dashboard.mjs).

    export_docx.py spec.json out.docx

spec.json: {"title": str, "subtitle": str, "parts": [{"title": str, "file": str, "width": int, "height": int}]}

Заголовки (документа и виджетов) — обычный текст Word, его можно править; виджет — картинка под ним.
"""
import json
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

# Виджеты на экране широкие, поэтому лист альбомный, а поля узкие.
PAGE_W, PAGE_H, MARGIN = Cm(29.7), Cm(21.0), Cm(1.2)


def main(spec_path, out_path):
    with open(spec_path, encoding='utf-8') as handle:
        spec = json.load(handle)

    document = Document()

    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = PAGE_W, PAGE_H
    section.left_margin = section.right_margin = MARGIN
    section.top_margin = section.bottom_margin = MARGIN

    normal = document.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(10)

    usable_w = PAGE_W - 2 * MARGIN
    usable_h = PAGE_H - 2 * MARGIN

    title = (spec.get('title') or '').strip()
    subtitle = (spec.get('subtitle') or '').strip()
    used_h = 0

    if title:
        heading = document.add_paragraph()
        run = heading.add_run(title)
        run.bold = True
        run.font.size = Pt(20)
        heading.paragraph_format.space_after = Pt(2)
        used_h += Cm(1.2)

    if subtitle:
        sub = document.add_paragraph()
        run = sub.add_run(subtitle)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x66, 0x70, 0x85)
        sub.paragraph_format.space_after = Pt(8)
        used_h += Cm(0.8)

    heading_h = Cm(1.0)

    for index, part in enumerate(spec.get('parts', [])):
        widget_title = (part.get('title') or '').strip()

        # Вписываем виджет в страницу целиком: картинку Word через границу листа не режет.
        ratio = part['width'] / part['height']
        # Первая страница уже занята заголовком документа.
        free_h = usable_h - (used_h if index == 0 else 0) - (heading_h if widget_title else 0) - Cm(0.4)
        # Чуть уже печатной области: так картинка не вылезает за поле, даже если редактор сдвинет её на пару пикселей.
        width = int(usable_w * 0.985)
        height = width / ratio

        if height > free_h:
            height = free_h
            width = height * ratio

        if widget_title:
            heading = document.add_paragraph()
            run = heading.add_run(widget_title)
            run.bold = True
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(0x1d, 0x27, 0x3b)
            heading.paragraph_format.space_before = Pt(6)
            heading.paragraph_format.space_after = Pt(4)
            # Заголовок не остаётся один в конце страницы без своего виджета.
            heading.paragraph_format.keep_with_next = True

        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.add_run().add_picture(part['file'], width=int(width), height=int(height))

    document.save(out_path)


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
